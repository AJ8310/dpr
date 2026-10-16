import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from dpr_engine.data_model import DPRDocumentModel
from dpr_engine.structures.structure_engine import DPRStructureEngine
from dpr_engine.utils.currency import format_inr

DOCX_I18N = {
    "en": {
        "vkf_title": "VISION KARNATAKA FOUNDATION\n",
        "dpr_title": "DETAILED PROJECT REPORT (DPR)",
        "track_prefix": "DPR Track",
        "depth_prefix": "Depth",
        "location_prefix": "Location",
        "ch1_heading": "1.0 Executive Summary & Project Highlights",
        "ch1_text": "{business_name} is a commercial enterprise located in {district}, {state}. The total project cost is estimated at {total_cost}, funded by promoter contribution of {promoter_contrib} and bank loan of {bank_loan}. The project features an average 5-year DSCR of {avg_dscr} and break-even point of {bep_percent}%.",
        "m_total_cost": "Total Project Cost",
        "m_promoter": "Promoter Equity Contribution",
        "m_loan": "Proposed Bank Loan",
        "m_subsidy": "Government Subsidy Claim",
        "m_dscr": "Average 5-Year DSCR",
        "ch2_heading": "2.0 Cost of Project & Means of Finance",
        "ch2_land": "Land & Site Development",
        "ch2_building": "Building & Civil Works",
        "ch2_furniture": "Furniture & Fixtures",
        "ch2_wc": "Working Capital Margin",
        "ch2_total": "Total Capital Outlay",
        "ch3_heading": "3.0 5-Year Projected Income Statement (P&L)",
        "pnl_headers": ["Metric", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"],
        "pnl_cap": "Capacity Utilization %",
        "pnl_rev": "Gross Revenue (₹)",
        "pnl_opex": "Total OPEX (₹)",
        "pnl_pat": "PAT (Net Profit) (₹)",
        "pnl_dscr": "DSCR Ratio",
        "ch4_heading": "4.0 Graphical Financial Projections & Analysis",
    },
    "kn": {
        "vkf_title": "ವಿಷನ್ ಕರ್ನಾಟಕ ಫೌಂಡೇಶನ್\n",
        "dpr_title": "ವಿವರವಾದ ಯೋಜನಾ ವರದಿ (DPR)",
        "track_prefix": "ಯೋಜನಾ ವಿಭಾಗ",
        "depth_prefix": "ವರದಿಯ ಆಳ",
        "location_prefix": "ಸ್ಥಳ",
        "ch1_heading": "1.0 ಕಾರ್ಯಕಾರಿ ಸಾರಾಂಶ ಮತ್ತು ಯೋಜನೆಯ ಮುಖ್ಯಾಂಶಗಳು",
        "ch1_text": "{business_name} ಸಂಸ್ಥೆಯು {district}, {state} ನಲ್ಲಿ ಸ್ಥಾಪಿಸಲಾಗುವ ವಾಣಿಜ್ಯ ಉದ್ಯಮವಾಗಿದೆ. ಯೋಜನೆಯ ಒಟ್ಟು ಅಂದಾಜು ವೆಚ್ಚ {total_cost} ಆಗಿದ್ದು, ಪ್ರವರ್ತಕರ ಇಕ್ವಿಟಿ {promoter_contrib} ಮತ್ತು ಬ್ಯಾಂಕ್ ಸಾಲ {bank_loan} ಮೂಲಕ ಹಣಕಾಸು ಒದಗಿಸಲಾಗಿದೆ. ಯೋಜನೆಯ ಸರಾಸರಿ 5-ವರ್ಷದ ಸಾಲ ಮರುಪಾವತಿ ಅನುಪಾತ (DSCR) {avg_dscr} ಮತ್ತು ಬ್ರೇಕ್-ಈವನ್ ಪಾಯಿಂಟ್ {bep_percent}% ಆಗಿದೆ.",
        "m_total_cost": "ಒಟ್ಟು ಯೋಜನಾ ವೆಚ್ಚ",
        "m_promoter": "ಪ್ರವರ್ತಕರ ಇಕ್ವಿಟಿ ಬಂಡವಾಳ",
        "m_loan": "ಪ್ರಸ್ತಾವಿತ ಬ್ಯಾಂಕ್ ಸಾಲ",
        "m_subsidy": "ಸರ್ಕಾರಿ ಸಬ್ಸಿಡಿ ಮೊತ್ತ",
        "m_dscr": "ಸರಾಸರಿ 5-ವರ್ಷದ DSCR ಅನುಪಾತ",
        "ch2_heading": "2.0 ಯೋಜನೆಯ ವೆಚ್ಚ ಮತ್ತು ಹಣಕಾಸಿನ ಮೂಲಗಳು",
        "ch2_land": "ಜಮೀನು ಮತ್ತು ಅಭಿವೃದ್ಧಿ",
        "ch2_building": "ಕಟ್ಟಡ ಮತ್ತು ಸಿವಿಲ್ ಕಾಮಗಾರಿ",
        "ch2_furniture": "ಪೀಠೋಪಕರಣಗಳು",
        "ch2_wc": "ದುಡಿಯುವ ಬಂಡವಾಳದ ಪಾಲು",
        "ch2_total": "ಒಟ್ಟು ಬಂಡವಾಳ ವೆಚ್ಚ",
        "ch3_heading": "3.0 5-ವರ್ಷಗಳ ಆದಾಯ ಮತ್ತು ಲಾಭದ ಅಂದಾಜು (P&L)",
        "pnl_headers": ["ವಿವರ", "ವರ್ಷ 1", "ವರ್ಷ 2", "ವರ್ಷ 3", "ವರ್ಷ 4", "ವರ್ಷ 5"],
        "pnl_cap": "ಸಾಮರ್ಥ್ಯ ಬಳಕೆ %",
        "pnl_rev": "ಒಟ್ಟು ವಹಿವಾಟು (₹)",
        "pnl_opex": "ಒಟ್ಟು ನಿರ್ವಹಣಾ ವೆಚ್ಚ (₹)",
        "pnl_pat": "ನಿವ್ವಳ ಲಾಭ (PAT) (₹)",
        "pnl_dscr": "DSCR ಅನುಪಾತ",
        "ch4_heading": "4.0 ಹಣಕಾಸಿನ ನಕ್ಷೆಗಳು ಮತ್ತು ವಿಶ್ಲೇಷಣೆ",
    },
    "hi": {
        "vkf_title": "विजन कर्नाटक फाउंडेशन\n",
        "dpr_title": "विस्तृत परियोजना रिपोर्ट (DPR)",
        "track_prefix": "डीपीआर श्रेणी",
        "depth_prefix": "रिपोर्ट गहराई",
        "location_prefix": "स्थान",
        "ch1_heading": "1.0 कार्यकारी सारांश एवं परियोजना मुख्य विशेषताएं",
        "ch1_text": "{business_name} एक वाणिज्यिक उद्यम है जो {district}, {state} में स्थापित है। परियोजना की कुल अनुमानित लागत {total_cost} है, जिसे प्रवर्तक अंशदान {promoter_contrib} एवं बैंक ऋण {bank_loan} द्वारा वित्तपोषित किया गया है। परियोजना का औसत 5-वर्षीय ऋण सेवा कवरेज अनुपात (DSCR) {avg_dscr} तथा ब्रेक-इवन बिंदु {bep_percent}% है।",
        "m_total_cost": "कुल परियोजना लागत",
        "m_promoter": "प्रवर्तक इक्विटी अंशदान",
        "m_loan": "प्रस्तावित बैंक ऋण",
        "m_subsidy": "सरकारी सब्सिडी दावा",
        "m_dscr": "औसत 5-वर्षीय DSCR",
        "ch2_heading": "2.0 परियोजना लागत एवं वित्तपोषण के साधन",
        "ch2_land": "भूमि एवं स्थल विकास",
        "ch2_building": "भवन एवं सिविल निर्माण",
        "ch2_furniture": "फर्नीचर एवं फिटिंग्स",
        "ch2_wc": "कार्यशील पूंजी मार्जिन",
        "ch2_total": "कुल पूंजी परिव्यय",
        "ch3_heading": "3.0 5-वर्षीय अनुमानित आय विवरण (लाभ एवं हानि)",
        "pnl_headers": ["मापदंड", "वर्ष 1", "वर्ष 2", "वर्ष 3", "वर्ष 4", "वर्ष 5"],
        "pnl_cap": "क्षमता उपयोग %",
        "pnl_rev": "सकल राजस्व (₹)",
        "pnl_opex": "कुल परिचालन व्यय (₹)",
        "pnl_pat": "शुद्ध लाभ (PAT) (₹)",
        "pnl_dscr": "DSCR अनुपात",
        "ch4_heading": "4.0 वित्तीय चार्ट एवं विश्लेषण",
    }
}

class DOCXDocumentRenderer:

    @classmethod
    def render_docx(cls, doc: DPRDocumentModel, output_docx_path: str) -> str:
        document = Document()
        lang = getattr(doc, 'language', 'en') or 'en'
        i18n = DOCX_I18N.get(lang, DOCX_I18N['en'])

        # Page Setup Margins & Page Borders
        for section in document.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

            pg_xml = (
                '<w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:top w:val="single" w:sz="12" w:space="24" w:color="0F172A"/>'
                '<w:left w:val="single" w:sz="12" w:space="24" w:color="0F172A"/>'
                '<w:bottom w:val="single" w:sz="12" w:space="24" w:color="0F172A"/>'
                '<w:right w:val="single" w:sz="12" w:space="24" w:color="0F172A"/>'
                '</w:pgBorders>'
            )
            section._sectPr.append(parse_xml(pg_xml))

        # Title / Cover Block
        title_p = document.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = title_p.add_run(i18n["vkf_title"])
        run_title.font.size = Pt(14)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        run_sub = title_p.add_run(f"{i18n['dpr_title']}\n{doc.business_name.upper()}\n")
        run_sub.font.size = Pt(18)
        run_sub.font.bold = True
        run_sub.font.color.rgb = RGBColor(30, 41, 59)

        run_track = title_p.add_run(f"{i18n['track_prefix']}: {doc.dpr_type} | {i18n['depth_prefix']}: {doc.dpr_depth.upper()}\n{i18n['location_prefix']}: {doc.district}, {doc.state}\n")
        run_track.font.size = Pt(10)
        run_track.font.color.rgb = RGBColor(100, 116, 139)

        document.add_heading(i18n["ch1_heading"], level=1)

        p_exec = document.add_paragraph()
        exec_text = i18n["ch1_text"].format(
            business_name=doc.business_name,
            district=doc.district,
            state=doc.state,
            total_cost=f"₹{format_inr(doc.total_cost)}",
            promoter_contrib=f"₹{format_inr(doc.promoter_contribution)}",
            bank_loan=f"₹{format_inr(doc.bank_loan)}",
            avg_dscr=doc.avg_dscr,
            bep_percent=doc.bep_percent
        )
        p_exec.add_run(exec_text)

        # Key Metrics Table
        table = document.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        metrics = [
            (i18n["m_total_cost"], f"₹{format_inr(doc.total_cost)}"),
            (i18n["m_promoter"], f"₹{format_inr(doc.promoter_contribution)}"),
            (i18n["m_loan"], f"₹{format_inr(doc.bank_loan)}"),
            (i18n["m_subsidy"], f"₹{format_inr(doc.subsidy)}"),
            (i18n["m_dscr"], f"{doc.avg_dscr}"),
        ]

        for i, (k, v) in enumerate(metrics):
            row = table.rows[i]
            row.cells[0].text = k
            row.cells[1].text = v

        document.add_heading(i18n["ch2_heading"], level=1)
        p_fin = document.add_paragraph()
        p_fin.add_run(
            f"{i18n['ch2_land']}: ₹{format_inr(doc.land_cost)}\n"
            f"{i18n['ch2_building']}: ₹{format_inr(doc.building_cost)}\n"
            f"{i18n['ch2_furniture']}: ₹{format_inr(doc.furniture_cost)}\n"
            f"{i18n['ch2_wc']}: ₹{format_inr(doc.working_capital)}\n"
            f"{i18n['ch2_total']}: ₹{format_inr(doc.total_cost)}"
        )

        # 5-Year Financial Projections Table
        if doc.projections:
            document.add_heading(i18n["ch3_heading"], level=1)
            p_table = document.add_table(rows=6, cols=6)
            p_table.style = 'Table Grid'

            headers = i18n["pnl_headers"]
            for col_idx, h in enumerate(headers):
                p_table.rows[0].cells[col_idx].text = h

            metrics_rows = [
                (i18n["pnl_cap"], [f"{p.capacity_utilization}%" for p in doc.projections]),
                (i18n["pnl_rev"], [f"₹{format_inr(p.revenue)}" for p in doc.projections]),
                (i18n["pnl_opex"], [f"₹{format_inr(p.total_opex)}" for p in doc.projections]),
                (i18n["pnl_pat"], [f"₹{format_inr(p.pat)}" for p in doc.projections]),
                (i18n["pnl_dscr"], [f"{p.dscr}" for p in doc.projections]),
            ]

            for row_idx, (m_label, m_vals) in enumerate(metrics_rows, start=1):
                r_cells = p_table.rows[row_idx].cells
                r_cells[0].text = m_label
                for c_idx, val in enumerate(m_vals, start=1):
                    r_cells[c_idx].text = val

        # Embed Generated Charts if available
        if doc.chart_paths:
            document.add_heading(i18n["ch4_heading"], level=1)
            for chart_key, c_path in doc.chart_paths.items():
                if os.path.exists(c_path):
                    document.add_paragraph(f"Chart: {chart_key.replace('_', ' ').title()}")
                    document.add_picture(c_path, width=Inches(5.5))

        document.save(output_docx_path)
        return output_docx_path
